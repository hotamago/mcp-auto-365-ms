"""Per-cell Excel writes through the Graph workbook API - request shapes offline.

The live PATCH is not exercised here (tests never touch the network), so these
tests pin down exactly what would go on the wire: which URL, which body, which
headers, in which order, and what happens when any of it fails.
"""

from __future__ import annotations

import io

import pytest
from mcp.server.mcpserver import MCPServer
from mcp.server.mcpserver.exceptions import ToolError
from openpyxl import Workbook, load_workbook

import tools as tools_mod
from common.errors import (
    AuthExpiredError,
    ConcurrentEditError,
    ConnectError,
    Mcp365Error,
    RateLimitedError,
    TransportError,
    UnsupportedOperationError,
)
from sharepoint import workbook

DRIVE = "b!drive"
ITEM = "01ITEM"
BASE = f"/drives/{DRIVE}/items/{ITEM}/workbook"
SHEET = "Sprint 4 (17.09)"
SHEET_ID = "{0E67021A-C793-4C5A-81DE-045870E88EAE}"
ENCODED_ID = "%7B0E67021A-C793-4C5A-81DE-045870E88EAE%7D"


def _http(status: int, message: str = "", cls: type[Mcp365Error] = Mcp365Error) -> Mcp365Error:
    """An error as ``common.http._perform`` raises it: typed, with ``http_status`` set."""
    err = cls(message or f"HTTP {status}")
    err.http_status = status
    return err


FORBIDDEN = _http(403, "Bị từ chối truy cập (HTTP 403)", AuthExpiredError)


class FakeGraph:
    """Records every call and replays canned answers."""

    def __init__(self, sheets=((SHEET, SHEET_ID),), fail=None, values=None, fail_patch=None):
        self.calls: list[dict] = []
        self._sheets = sheets
        self._fail = fail or {}
        self._fail_patch = fail_patch
        self._values = values or {}

    def __call__(self, path, method="GET", body=None, session_id="", context=""):
        self.calls.append(
            {"path": path, "method": method, "body": body, "session_id": session_id, "context": context}
        )
        if self._fail_patch is not None and method == "PATCH":
            raise self._fail_patch
        for marker, exc in self._fail.items():
            if marker in path:
                raise exc
        if "/worksheets?" in path:
            return {"value": [{"name": n, "id": i} for n, i in self._sheets]}
        if path.endswith("/createSession"):
            return {"id": "SESSION-1", "persistChanges": True}
        if path.endswith("/closeSession"):
            return {}
        if path.endswith("/worksheets/add"):
            return {"id": "{NEW-SHEET}", "name": body.get("name")}
        if "/range(address=" in path:
            ref = path.split("address='")[1].split("'")[0]
            if method == "PATCH":
                return {"address": f"'{SHEET}'!{ref}", "values": body["values"]}
            return {"address": f"'{SHEET}'!{ref}", "values": [[self._values.get(ref, "")]]}
        raise AssertionError(f"unexpected call {method} {path}")

    def paths(self, method=None):
        return [c["path"] for c in self.calls if method is None or c["method"] == method]


# ------------------------------------------------------------------ addresses


def test_addresses_are_normalised():
    assert workbook.check_addresses({" q34 ": "a", "$B$7": "b"}) == {"Q34": "a", "B7": "b"}


@pytest.mark.parametrize("bad", ["A1:B2", "A:A", "1", "not-a-cell", "A0", ""])
def test_only_single_cells_are_accepted(bad):
    with pytest.raises(Mcp365Error):
        workbook.check_addresses({bad: "x"})


# ----------------------------------------------------------------- fallbacks


def test_xlsm_falls_back_to_the_whole_file_path():
    with pytest.raises(workbook.WorkbookUnsupported) as excinfo:
        workbook.plan(FakeGraph(), DRIVE, ITEM, SHEET, {"A1": "x"}, "macro.xlsm")
    assert ".xlsx" in excinfo.value.reason


def test_cloning_a_sheet_falls_back():
    with pytest.raises(workbook.WorkbookUnsupported) as excinfo:
        workbook.plan(FakeGraph(), DRIVE, ITEM, "NEW", {"A1": "x"}, "plan.xlsx", copy_sheet_from="SYS2_PIN")
    assert "copy_sheet_from" in excinfo.value.reason


def test_workbook_api_unavailable_falls_back():
    """A Graph 404/403 on the sheet listing means the API cannot serve this file."""
    graph = FakeGraph(fail={"/worksheets?": FORBIDDEN})
    with pytest.raises(workbook.WorkbookUnsupported) as excinfo:
        workbook.plan(graph, DRIVE, ITEM, SHEET, {"A1": "x"}, "plan.xlsx")
    assert "403" in excinfo.value.reason


def test_a_locked_cell_read_is_an_error_not_a_fallback():
    """423 means someone holds the file - the whole-file PUT would lose to them too."""
    graph = FakeGraph(fail={"range(address='Q34')": _http(423, "HTTP 423 Locked", ConcurrentEditError)})
    with pytest.raises(Mcp365Error) as excinfo:
        workbook.plan(graph, DRIVE, ITEM, SHEET, {"Q34": "x"}, "plan.xlsx")
    assert "KHÔNG chuyển sang ghi đè" in str(excinfo.value)


@pytest.mark.parametrize(
    "error",
    [
        TransportError("Request timed out after 30s."),
        ConnectError("Không kết nối được."),
        _http(503, "HTTP 503", RateLimitedError),
        _http(429, "HTTP 429", RateLimitedError),
        _http(500, "HTTP 500 Internal Server Error"),
    ],
    ids=["timeout", "connect", "503", "429", "500"],
)
def test_a_transient_failure_listing_sheets_does_not_fall_back(error):
    with pytest.raises(Mcp365Error) as excinfo:
        workbook.plan(FakeGraph(fail={"/worksheets?": error}), DRIVE, ITEM, SHEET, {"A1": "x"}, "plan.xlsx")
    assert "Chưa ghi gì" in str(excinfo.value)


@pytest.mark.parametrize(
    ("error", "refused"),
    [
        (_http(403, cls=AuthExpiredError), True),
        (_http(404, "HTTP 404 itemNotFound"), True),
        (_http(401, cls=AuthExpiredError), True),
        (_http(501, "HTTP 501 Not Implemented"), True),
        (_http(400, 'HTTP 400. Phản hồi: {"error":{"code":"NotSupported"}}'), True),
        (_http(400, 'HTTP 400. Phản hồi: {"error":{"code":"InvalidArgument"}}'), False),
        (UnsupportedOperationError("Không tìm thấy Azure CLI (az)."), True),
        (AuthExpiredError("az account get-access-token thất bại."), True),
        (_http(409, cls=ConcurrentEditError), False),
        (_http(412, cls=ConcurrentEditError), False),
        (_http(423, cls=ConcurrentEditError), False),
        (_http(502, "HTTP 502 Bad Gateway"), False),
        (RateLimitedError("Đã thử lại 3 lần nhưng vẫn thất bại"), False),
        (TransportError("timed out"), False),
        (Mcp365Error("Azure CLI không phản hồi sau 60s."), False),
    ],
)
def test_only_a_definitive_refusal_counts_as_one(error, refused):
    assert workbook.graph_refused(error) is refused


def test_no_cells_is_an_error_not_a_fallback():
    with pytest.raises(Mcp365Error):
        workbook.plan(FakeGraph(), DRIVE, ITEM, SHEET, {}, "plan.xlsx")


# --------------------------------------------------------------- preview/plan


def test_plan_resolves_the_sheet_by_id_and_only_reads():
    """The id goes in the URL bare and percent-encoded: the quoted form 404s."""
    graph = FakeGraph(values={"Q34": "cũ"})
    sheet_id, changes = workbook.plan(graph, DRIVE, ITEM, SHEET, {"Q34": "mới"}, "plan.xlsx")
    assert sheet_id == SHEET_ID
    assert changes == [{"cell": "Q34", "old": "cũ", "new": "mới"}]
    assert graph.paths() == [
        f"{BASE}/worksheets?$select=id,name",
        f"{BASE}/worksheets/{ENCODED_ID}/range(address='Q34')?$select=values",
    ]
    assert all(c["method"] == "GET" for c in graph.calls), "preview must not write"
    assert not any(c["session_id"] for c in graph.calls), "preview must not open a session"


def test_plan_for_a_missing_sheet_reports_it_will_be_created():
    graph = FakeGraph()
    sheet_id, changes = workbook.plan(graph, DRIVE, ITEM, "Sprint 5", {"A1": "x"}, "plan.xlsx")
    assert sheet_id == ""
    assert changes[0]["cell"] == "(sheet)" and "Sprint 5" in changes[0]["new"]


# ---------------------------------------------------------------------- write


def test_apply_patches_each_cell_inside_one_session():
    graph = FakeGraph()
    written, warnings = workbook.apply(graph, DRIVE, ITEM, SHEET, SHEET_ID, {"Q34": "xong", "R34": 7})
    assert (written, warnings) == (2, [])

    assert graph.paths("POST") == [f"{BASE}/createSession", f"{BASE}/closeSession"]
    patches = [c for c in graph.calls if c["method"] == "PATCH"]
    assert [c["path"] for c in patches] == [
        f"{BASE}/worksheets/{ENCODED_ID}/range(address='Q34')",
        f"{BASE}/worksheets/{ENCODED_ID}/range(address='R34')",
    ]
    assert [c["body"] for c in patches] == [{"values": [["xong"]]}, {"values": [[7]]}]


def test_the_session_id_travels_on_every_call_after_it_opens():
    graph = FakeGraph()
    workbook.apply(graph, DRIVE, ITEM, SHEET, SHEET_ID, {"Q34": "x"})
    after_open = graph.calls[1:]  # everything past createSession
    assert after_open and all(c["session_id"] == "SESSION-1" for c in after_open)
    assert graph.calls[0]["body"] == {"persistChanges": True}


def test_the_session_is_closed_even_when_a_patch_blows_up():
    graph = FakeGraph(fail_patch=_http(500, "HTTP 500"))
    with pytest.raises(Mcp365Error) as excinfo:
        workbook.apply(graph, DRIVE, ITEM, SHEET, SHEET_ID, {"Q34": "x"})
    assert not isinstance(excinfo.value, workbook.WorkbookUnsupported)
    assert f"{BASE}/closeSession" in graph.paths("POST")


def test_a_refused_first_patch_asks_for_the_whole_file_path():
    """The likely live outcome if the Azure CLI token cannot write: reads work,
    the first PATCH 403s. A dead tool would be worse than a disclosed downgrade."""
    graph = FakeGraph(fail_patch=FORBIDDEN)
    with pytest.raises(workbook.WorkbookUnsupported) as excinfo:
        workbook.apply(graph, DRIVE, ITEM, SHEET, SHEET_ID, {"Q34": "x"})
    assert "403" in excinfo.value.reason


@pytest.mark.parametrize(
    "error",
    [
        ConnectError("Không kết nối được."),
        _http(503, "HTTP 503", RateLimitedError),
        _http(429, "HTTP 429", RateLimitedError),
        _http(502, "HTTP 502 Bad Gateway"),
        _http(423, "HTTP 423 Locked", ConcurrentEditError),
        _http(412, "HTTP 412", ConcurrentEditError),
    ],
    ids=["connect", "503", "429", "502", "423", "412"],
)
def test_a_transient_first_patch_failure_never_falls_back(error):
    graph = FakeGraph(fail_patch=error)
    with pytest.raises(Mcp365Error) as excinfo:
        workbook.apply(graph, DRIVE, ITEM, SHEET, SHEET_ID, {"Q34": "x"})
    assert not isinstance(excinfo.value, workbook.WorkbookUnsupported)
    assert "KHÔNG chuyển sang ghi đè" in str(excinfo.value)


def test_a_lost_reply_on_the_first_patch_says_the_cell_may_have_been_written():
    """The PATCH was sent; Graph may have applied it. Overwriting the file now would
    race our own write - stop and have the user look at the cell."""
    graph = FakeGraph(fail_patch=TransportError("Request timed out after 120s."))
    with pytest.raises(Mcp365Error) as excinfo:
        workbook.apply(graph, DRIVE, ITEM, SHEET, SHEET_ID, {"Q34": "x", "R34": "y"})
    assert not isinstance(excinfo.value, workbook.WorkbookUnsupported)
    text = str(excinfo.value)
    assert "`Q34`" in text and "có thể ĐÃ được ghi" in text
    assert len(graph.paths("PATCH")) == 1, "must stop at the unsure cell"


def test_a_created_sheet_counts_as_written_so_a_failed_patch_never_falls_back():
    """`worksheets/add` changed the file: a whole-file PUT now would be a second writer."""
    graph = FakeGraph(fail_patch=FORBIDDEN)
    with pytest.raises(Mcp365Error) as excinfo:
        workbook.apply(graph, DRIVE, ITEM, "Sprint 5", "", {"A1": "x", "B1": "y"})
    assert not isinstance(excinfo.value, workbook.WorkbookUnsupported)
    assert "Đã tạo sheet 'Sprint 5', ghi được 0/2 ô" in str(excinfo.value)


def test_a_refused_sheet_creation_still_falls_back():
    graph = FakeGraph(fail={"worksheets/add": FORBIDDEN})
    with pytest.raises(workbook.WorkbookUnsupported):
        workbook.apply(graph, DRIVE, ITEM, "Sprint 5", "", {"A1": "x"})


def test_a_lost_reply_creating_a_sheet_says_it_may_exist():
    graph = FakeGraph(fail={"worksheets/add": TransportError("timed out")})
    with pytest.raises(Mcp365Error) as excinfo:
        workbook.apply(graph, DRIVE, ITEM, "Sprint 5", "", {"A1": "x"})
    assert not isinstance(excinfo.value, workbook.WorkbookUnsupported)
    assert "có thể ĐÃ được tạo" in str(excinfo.value)
    assert not graph.paths("PATCH")


def test_a_refused_session_still_writes_sessionless():
    """The session is an optimisation, not a requirement."""
    graph = FakeGraph(fail={"createSession": Mcp365Error("HTTP 403 Forbidden")})
    written, warnings = workbook.apply(graph, DRIVE, ITEM, SHEET, SHEET_ID, {"Q34": "x"})
    assert written == 1
    assert warnings and "403" in warnings[0]
    assert not any(c["session_id"] for c in graph.calls)


def test_a_failed_close_never_fails_the_write():
    graph = FakeGraph(fail={"closeSession": Mcp365Error("HTTP 500")})
    written, warnings = workbook.apply(graph, DRIVE, ITEM, SHEET, SHEET_ID, {"Q34": "x"})
    assert written == 1 and any("hết hạn" in w for w in warnings)


def test_a_partial_write_says_what_landed_and_is_never_retried_whole_file():
    """Re-uploading the workbook after a cell landed would undo the co-authoring."""
    graph = FakeGraph()
    patched: list[str] = []

    def second_patch_fails(path, method="GET", body=None, session_id="", context=""):
        if method == "PATCH":
            patched.append(path)
            if len(patched) == 2:
                raise _http(423, "HTTP 423 Locked", ConcurrentEditError)
        return graph(path, method, body, session_id, context)

    with pytest.raises(Mcp365Error) as excinfo:
        workbook.apply(second_patch_fails, DRIVE, ITEM, SHEET, SHEET_ID, {"Q34": "a", "R34": "b"})
    assert "1/2" in str(excinfo.value)
    # An Mcp365Error, not a WorkbookUnsupported: the caller must not fall back.
    assert not isinstance(excinfo.value, workbook.WorkbookUnsupported)


def test_a_missing_sheet_is_created_before_the_first_patch():
    graph = FakeGraph()
    written, _ = workbook.apply(graph, DRIVE, ITEM, "Sprint 5", "", {"A1": "x"})
    assert written == 1
    assert graph.paths("POST")[0] == f"{BASE}/worksheets/add"
    assert graph.calls[0]["body"] == {"name": "Sprint 5"}
    assert "/worksheets/%7BNEW-SHEET%7D/range(address='A1')" in graph.paths("PATCH")[0]


def test_a_value_excel_stored_differently_is_reported():
    class Coercing(FakeGraph):
        def __call__(self, path, method="GET", body=None, session_id="", context=""):
            res = super().__call__(path, method, body, session_id, context)
            if method == "PATCH":
                return {**res, "values": [["KHÁC"]]}
            return res

    _written, warnings = workbook.apply(Coercing(), DRIVE, ITEM, SHEET, SHEET_ID, {"Q34": "gửi"})
    assert warnings and "KHÁC" in warnings[0]


def test_render_changes_shows_empty_cells_explicitly():
    out = workbook.render_changes([{"cell": "Q34", "old": "", "new": "xong"}], SHEET)
    assert "`Q34`" in out and "_(trống)_" in out


# ------------------------------------------- which path `update_sharepoint_sheet` takes


def _xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["Q34"] = "cũ"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class FakeSharePoint:
    def __init__(self, name="plan.xlsx", graph=None):
        self.name = name
        self.graph = graph or FakeGraph(values={"Q34": "cũ"})
        self.uploaded: bytes | None = None
        self.downloads = 0

    def resolve_file(self, _ref):
        return DRIVE, {"id": ITEM, "name": self.name, "eTag": "etag-1", "webUrl": "https://sp/plan.xlsx"}

    def call_workbook(self, path, method="GET", body=None, session_id="", context=""):
        return self.graph(path, method, body, session_id, context)

    def read_file_bytes(self, _drive_id, _item):
        self.downloads += 1
        return _xlsx_bytes()

    def put_file_bytes(self, _drive_id, _item_id, data, if_match=""):
        self.uploaded = data
        assert if_match == "etag-1"
        return {"webUrl": "https://sp/plan.xlsx"}


def _server(monkeypatch, fake):
    monkeypatch.setattr(tools_mod, "sp", lambda: fake)
    mcp = MCPServer("t")
    tools_mod.register_all(mcp)
    return mcp


@pytest.mark.anyio
async def test_an_xlsx_is_written_cell_by_cell_and_never_re_uploaded(monkeypatch):
    """The whole point: no download, no whole-file PUT, so an open file is fine."""
    fake = FakeSharePoint()
    mcp = _server(monkeypatch, fake)
    await mcp.call_tool(
        "update_sharepoint_sheet",
        {"file_url_or_guid": "GUID", "sheet": SHEET, "cells": {"Q34": "xong"}, "is_user_confirm": True},
    )
    assert fake.downloads == 0 and fake.uploaded is None
    assert [c["method"] for c in fake.graph.calls if c["method"] == "PATCH"] == ["PATCH"]


@pytest.mark.anyio
async def test_the_preview_call_writes_nothing_and_shows_the_change_list(monkeypatch):
    fake = FakeSharePoint()
    mcp = _server(monkeypatch, fake)
    with pytest.raises(ToolError) as excinfo:
        await mcp.call_tool(
            "update_sharepoint_sheet",
            {"file_url_or_guid": "GUID", "sheet": SHEET, "cells": {"Q34": "xong"}, "is_user_confirm": False},
        )
    refusal = str(excinfo.value)
    assert "`Q34`" in refusal and "cũ" in refusal and "xong" in refusal
    assert not [c for c in fake.graph.calls if c["method"] != "GET"]
    assert fake.uploaded is None


@pytest.mark.anyio
async def test_cloning_a_sheet_falls_back_to_the_whole_file_path_and_says_why(monkeypatch):
    fake = FakeSharePoint()
    mcp = _server(monkeypatch, fake)
    res = await mcp.call_tool(
        "update_sharepoint_sheet",
        {
            "file_url_or_guid": "GUID",
            "sheet": "Sprint 5",
            "cells": {"A1": "mới"},
            "copy_sheet_from": SHEET,
            "is_user_confirm": True,
        },
    )
    assert fake.downloads == 1 and fake.uploaded is not None
    assert load_workbook(io.BytesIO(fake.uploaded))["Sprint 5"]["A1"].value == "mới"
    assert "copy_sheet_from" in str(res.content)


@pytest.mark.anyio
async def test_the_fallback_preview_warns_it_overwrites_the_whole_file(monkeypatch):
    """A silent downgrade would look like a bug the next time a write hits 423."""
    fake = FakeSharePoint()
    mcp = _server(monkeypatch, fake)
    with pytest.raises(ToolError) as excinfo:
        await mcp.call_tool(
            "update_sharepoint_sheet",
            {
                "file_url_or_guid": "GUID",
                "sheet": "Sprint 5",
                "cells": {"A1": "mới"},
                "copy_sheet_from": SHEET,
                "is_user_confirm": False,
            },
        )
    refusal = str(excinfo.value)
    assert "ghi đè cả file" in refusal and "copy_sheet_from" in refusal
    assert fake.uploaded is None


@pytest.mark.anyio
async def test_graph_refusing_before_any_write_falls_back_instead_of_failing(monkeypatch):
    fake = FakeSharePoint(graph=FakeGraph(fail={"/worksheets?": FORBIDDEN}))
    mcp = _server(monkeypatch, fake)
    res = await mcp.call_tool(
        "update_sharepoint_sheet",
        {"file_url_or_guid": "GUID", "sheet": SHEET, "cells": {"Q34": "xong"}, "is_user_confirm": True},
    )
    assert fake.uploaded is not None
    assert "403" in str(res.content)


@pytest.mark.anyio
async def test_graph_refusing_the_first_patch_falls_back_and_still_writes(monkeypatch):
    """The expected live failure if the Azure CLI token cannot write cells."""
    fake = FakeSharePoint(graph=FakeGraph(values={"Q34": "cũ"}, fail_patch=FORBIDDEN))
    mcp = _server(monkeypatch, fake)
    res = await mcp.call_tool(
        "update_sharepoint_sheet",
        {"file_url_or_guid": "GUID", "sheet": SHEET, "cells": {"Q34": "xong"}, "is_user_confirm": True},
    )
    assert fake.uploaded is not None, "the edit must still land via the whole-file path"
    assert load_workbook(io.BytesIO(fake.uploaded))[SHEET]["Q34"].value == "xong"
    assert "403" in str(res.content)
    # The session opened for the attempt is still closed.
    assert f"{BASE}/closeSession" in fake.graph.paths("POST")


@pytest.mark.anyio
@pytest.mark.parametrize(
    "error",
    [TransportError("timed out"), _http(503, "HTTP 503", RateLimitedError), _http(429, "HTTP 429", RateLimitedError)],
    ids=["timeout", "503", "429"],
)
async def test_a_transient_first_patch_failure_never_overwrites_the_file(monkeypatch, error):
    fake = FakeSharePoint(graph=FakeGraph(values={"Q34": "cũ"}, fail_patch=error))
    mcp = _server(monkeypatch, fake)
    with pytest.raises(ToolError):
        await mcp.call_tool(
            "update_sharepoint_sheet",
            {"file_url_or_guid": "GUID", "sheet": SHEET, "cells": {"Q34": "xong"}, "is_user_confirm": True},
        )
    assert fake.uploaded is None, "a transient error must never turn into a whole-file PUT"


@pytest.mark.anyio
async def test_the_per_cell_draft_discloses_the_possible_downgrade(monkeypatch):
    """Approving the change list also approves the fallback, so it must be shown."""
    fake = FakeSharePoint()
    mcp = _server(monkeypatch, fake)
    with pytest.raises(ToolError) as excinfo:
        await mcp.call_tool(
            "update_sharepoint_sheet",
            {"file_url_or_guid": "GUID", "sheet": SHEET, "cells": {"Q34": "xong"}, "is_user_confirm": False},
        )
    assert "ghi đè cả file" in str(excinfo.value)
