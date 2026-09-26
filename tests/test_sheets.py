"""In-memory workbook reading and editing: no network involved."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from common.errors import Mcp365Error
from sharepoint import sheets


def _checklist() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "SYS2_PIN"
    ws.append(["Feature", "Battery"])
    ws.append(["ID", "Review Item", "Description", "Yes/No", "Comment"])
    ws.append(["SYS2-001", "Complete Description", "Mô tả đầy đủ", "Yes", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_render_lists_sheets_when_none_is_named():
    out = sheets.render_sheet(_checklist(), name="c.xlsx")
    assert "`SYS2_PIN`" in out


def test_render_shows_a1_addresses():
    out = sheets.render_sheet(_checklist(), sheet="SYS2_PIN", name="c.xlsx")
    assert "| # | A | B | C | D | E |" in out
    assert "| 3 | SYS2-001 |" in out


def test_render_unknown_sheet_lists_the_real_ones():
    with pytest.raises(Mcp365Error) as excinfo:
        sheets.render_sheet(_checklist(), sheet="nope")
    assert "SYS2_PIN" in excinfo.value.remediation


def test_non_xlsx_bytes_explain_the_likely_cause():
    """A 'downloaded' xlsx is sometimes an HTML login page - say so."""
    with pytest.raises(Mcp365Error) as excinfo:
        sheets.render_sheet(b"<html>sign in</html>", name="x.xlsx")
    assert "HTML" in excinfo.value.remediation


def test_http_412_becomes_a_concurrent_edit_error():
    import urllib.error

    from common.errors import ConcurrentEditError, classify_http_error

    exc = urllib.error.HTTPError("https://graph", 412, "Precondition Failed", {}, io.BytesIO(b"{}"))
    assert isinstance(classify_http_error(exc, "ghi"), ConcurrentEditError)


# ---------------------------------------------------------------- merged cells


def _merged() -> bytes:
    """A sheet shaped like the real plan workbook: merged banners on row 1-2.

    ``ViTa - S5 - Management Plan.xlsx`` crashed ``render_sheet`` because row 1
    column B is a ``MergedCell``, which has no ``column_letter``.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sprint 4 (17.09)"
    ws.append(["Tracking", None, None, "Overall"])
    ws.append(["Sprint 4: 17/09", None, None, None])
    ws.append(["ID", "Task", "Owner", "Status"])
    ws.merge_cells("A1:C1")
    ws.merge_cells("A2:D2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_render_survives_merged_cells():
    out = sheets.render_sheet(_merged(), sheet="Sprint 4 (17.09)", name="plan.xlsx")
    assert "| # | A | B | C | D |" in out


def test_merged_cells_show_the_anchor_value_and_blank_followers():
    out = sheets.render_sheet(_merged(), sheet="Sprint 4 (17.09)", name="plan.xlsx")
    assert "| 1 | Tracking |  |  | Overall |" in out


def test_render_lists_the_merged_ranges():
    """The reader sees which cells are merged and where the value lives."""
    out = sheets.render_sheet(_merged(), sheet="Sprint 4 (17.09)", name="plan.xlsx")
    assert "`A1:C1`" in out and "`A2:D2`" in out


def test_a1_addresses_stay_aligned_with_the_merged_sheet():
    """Column letters and row numbers printed must still address the real cells."""
    out = sheets.render_sheet(_merged(), sheet="Sprint 4 (17.09)", name="plan.xlsx")
    assert "| 3 | ID | Task | Owner | Status |" in out


def test_a_whole_row_merge_does_not_blow_up():
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "banner"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=16384)
    buf = io.BytesIO()
    wb.save(buf)
    out = sheets.render_sheet(buf.getvalue(), sheet="S", name="wide.xlsx")
    assert "Ô gộp (1)" in out


def test_a_sheet_name_matches_ignoring_case_and_surrounding_spaces():
    """Real names carry trailing spaces (`'S5 Feature Release Plan '`)."""
    assert "SYS2-001" in sheets.render_sheet(_checklist(), sheet=" sys2_pin ")


# ------------------------------------------------------------ hidden sheets
# 26/09: a BA workbook showed 3 tabs and hid 12; data from a hidden tab was
# taken as the source of truth. Hidden tabs are hidden on purpose.


def _with_hidden_tabs() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "UC List"
    ws.append(["UC", "Status"])
    ws.append(["UC01", "Approved"])
    old = wb.create_sheet("Old Draft")
    old.append(["UC01", "SECRET-DRAFT"])
    old.sheet_state = "hidden"
    lookup = wb.create_sheet("Lookup")
    lookup.append(["VERY-HIDDEN-VALUE"])
    lookup.sheet_state = "veryHidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_listing_names_hidden_sheets_but_skips_them():
    out = sheets.render_sheet(_with_hidden_tabs(), name="wmc.xlsx")
    assert "| `UC List` | 2 × 2 |" in out
    assert "| `Old Draft` | (ẩn, bỏ qua) |" in out
    assert "| `Lookup` | (ẩn, bỏ qua) |" in out  # veryHidden counts as hidden
    assert "1 hiện, 2 ẩn" in out
    assert "include_hidden=True" in out


def test_listing_with_include_hidden_shows_their_size_and_marks_them():
    out = sheets.render_sheet(_with_hidden_tabs(), name="wmc.xlsx", include_hidden=True)
    assert "| `Old Draft` (ẩn) | 1 × 2 |" in out
    assert "| `Lookup` (ẩn) | 1 × 1 |" in out
    assert "bỏ qua" not in out


def test_visible_sheet_reads_as_before():
    out = sheets.render_sheet(_with_hidden_tabs(), sheet="UC List", name="wmc.xlsx")
    assert "| 2 | UC01 | Approved |" in out
    assert "sheet ẩn" not in out


@pytest.mark.parametrize("sheet, state", [("Old Draft", "hidden"), ("lookup", "veryHidden")])
def test_naming_a_hidden_sheet_is_refused_clearly(sheet, state):
    with pytest.raises(Mcp365Error) as excinfo:
        sheets.render_sheet(_with_hidden_tabs(), sheet=sheet, name="wmc.xlsx")
    err = excinfo.value
    assert "bị ẩn" in err.message and state in err.message
    assert "SECRET-DRAFT" not in err.message and "VERY-HIDDEN-VALUE" not in err.message
    assert "include_hidden=True" in err.remediation
    assert "UC List" in err.remediation


@pytest.mark.parametrize("sheet, value", [("Old Draft", "SECRET-DRAFT"), ("Lookup", "VERY-HIDDEN-VALUE")])
def test_include_hidden_reads_a_hidden_sheet_and_says_so(sheet, value):
    out = sheets.render_sheet(_with_hidden_tabs(), sheet=sheet, name="wmc.xlsx", include_hidden=True)
    assert value in out
    assert "— sheet ẩn" in out


def test_unknown_sheet_suggests_only_visible_ones_by_default():
    with pytest.raises(Mcp365Error) as excinfo:
        sheets.render_sheet(_with_hidden_tabs(), sheet="nope")
    assert "UC List" in excinfo.value.remediation
    assert "Old Draft" not in excinfo.value.remediation


@pytest.mark.anyio
async def test_read_sharepoint_sheet_tool_passes_include_hidden_through(monkeypatch):
    from mcp.server.mcpserver import MCPServer
    from mcp.server.mcpserver.exceptions import ToolError

    import tools as tools_mod

    class FakeSharePoint:
        def resolve_file(self, _ref):
            return "drv", {"id": "1", "name": "wmc.xlsx"}

        def read_file_bytes(self, _drive_id, _item):
            return _with_hidden_tabs()

    monkeypatch.setattr(tools_mod, "sp", lambda: FakeSharePoint())
    mcp = MCPServer("t")
    tools_mod.register_all(mcp)

    listing = await mcp.call_tool("read_sharepoint_sheet", {"file_url_or_guid": "GUID"})
    assert "(ẩn, bỏ qua)" in str(listing.content)
    with pytest.raises(ToolError) as excinfo:
        await mcp.call_tool("read_sharepoint_sheet", {"file_url_or_guid": "GUID", "sheet": "Old Draft"})
    assert "include_hidden=True" in str(excinfo.value) and "SECRET-DRAFT" not in str(excinfo.value)
    res = await mcp.call_tool(
        "read_sharepoint_sheet", {"file_url_or_guid": "GUID", "sheet": "Old Draft", "include_hidden": True}
    )
    assert "SECRET-DRAFT" in str(res.content)
