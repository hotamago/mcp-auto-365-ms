"""Read ``.xlsx`` workbooks in memory.

Pure functions over bytes: no network, no credentials. The SharePoint client
fetches; this module only turns bytes into a table. There is no write path: the
tools that edited existing files (cells, Word comments, whole-file replace) were
removed on 26/09 - writes into a file others had open were refused (423/412) or
had to re-upload the whole file.

**Merged cells.** A merged range stores its value in the top-left (anchor) cell
only; every other cell of the range is an openpyxl ``MergedCell`` with a
read-only ``None`` value. ``render_sheet`` follows that model rather than
spreading the anchor's value: it prints the anchor's value at the anchor address,
leaves the rest blank and lists the ranges under the table.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from common.errors import Mcp365Error


def _open(data: bytes, name: str = ""):
    try:
        return load_workbook(io.BytesIO(data), keep_vba=name.lower().endswith(".xlsm"))
    except Exception as exc:
        raise Mcp365Error(
            f"Không mở được file Excel{f' `{name}`' if name else ''}: {exc}",
            "File có thể không phải .xlsx/.xlsm thật (VD: bản tải về là trang HTML đăng nhập). "
            "Tải lại hoặc mở bằng Excel Online để kiểm tra.",
        ) from exc


# How many merged ranges the footer lists before summarising the rest.
_MERGED_PREVIEW = 12


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("|", "\\|").replace("\n", " ")


def _merged_ranges(ws) -> list[Any]:
    """The sheet's merged ranges, top-left first, as a stable list."""
    return sorted(ws.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col))


def find_sheet(names: list[str], sheet: str) -> str:
    """The workbook's own name for ``sheet``: exact match first, then ignoring case
    and surrounding spaces (real names carry trailing spaces: ``'S5 Feature Release Plan '``).
    Returns ``""`` when there is none."""
    if sheet in names:
        return sheet
    wanted = sheet.casefold().strip()
    return next((n for n in names if n.casefold().strip() == wanted), "")


def is_hidden(ws) -> bool:
    """True for a sheet the author hid: ``hidden`` (Unhide… in Excel) or ``veryHidden`` (VBA only)."""
    return getattr(ws, "sheet_state", "visible") != "visible"


def _list_sheets(wb, name: str, include_hidden: bool) -> str:
    hidden = [ws.title for ws in wb.worksheets if is_hidden(ws)]
    head = f"# Workbook `{name}` — {len(wb.sheetnames)} sheet"
    if hidden:
        head += f" ({len(wb.worksheets) - len(hidden)} hiện, {len(hidden)} ẩn)"
    rows = [head + "\n", "| Sheet | Kích thước |", "| --- | --- |"]
    for ws in wb.worksheets:
        if not is_hidden(ws):
            rows.append(f"| `{ws.title}` | {ws.max_row} × {ws.max_column} |")
        elif include_hidden:
            rows.append(f"| `{ws.title}` (ẩn) | {ws.max_row} × {ws.max_column} |")
        else:
            rows.append(f"| `{ws.title}` | (ẩn, bỏ qua) |")
    if hidden and not include_hidden:
        rows.append(
            f"\n_{len(hidden)} sheet ẩn không được đọc: tác giả file chủ động ẩn, đừng lấy làm căn cứ. "
            "Chỉ khi thật sự cần mới gọi lại với `include_hidden=True`._"
        )
    return "\n".join(rows)


def render_sheet(data: bytes, sheet: str = "", max_rows: int = 60, name: str = "", include_hidden: bool = False) -> str:
    """List the sheets, or render one sheet as a Markdown table with A1 refs.

    Hidden sheets (``hidden`` and ``veryHidden``) are skipped unless
    ``include_hidden``: the listing names them as "(ẩn, bỏ qua)", and naming one
    raises a clear error instead of rendering it. A BA hides a tab on purpose;
    reading one as the source of truth was a real mistake (26/09).

    Merged ranges keep openpyxl's model: only the top-left (anchor) cell carries
    the value, the rest of the range prints blank. The ranges are listed under
    the table.
    """
    wb = _open(data, name)
    if not sheet:
        return _list_sheets(wb, name, include_hidden)

    actual = find_sheet(wb.sheetnames, sheet)
    visible = [ws.title for ws in wb.worksheets if not is_hidden(ws)]
    if not actual:
        shown = wb.sheetnames if include_hidden else visible
        raise Mcp365Error(f"Không có sheet '{sheet}'.", f"Các sheet hiện có: {', '.join(shown)}")
    ws = wb[actual]
    if is_hidden(ws) and not include_hidden:
        raise Mcp365Error(
            f"Sheet '{actual}' đang bị ẩn trong file ({ws.sheet_state}), nên không được đọc. "
            "Tác giả file chủ động ẩn tab này; đừng lấy nó làm căn cứ.",
            "Chỉ khi thật sự cần mới gọi lại với `include_hidden=True`. "
            f"Các sheet đang hiện: {', '.join(visible) or '(không có)'}",
        )
    width = ws.max_column
    # ``cell().column_letter`` blows up on a MergedCell (it has no address of its
    # own), so the header is built from the column index instead.
    letters = [get_column_letter(c) for c in range(1, width + 1)]
    note = " — sheet ẩn" if is_hidden(ws) else ""
    out = [f"# `{name}` › `{sheet}` ({ws.max_row} × {width}){note}\n", "| # | " + " | ".join(letters) + " |"]
    out.append("| --- " * (width + 1) + "|")
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows), values_only=True), start=1):
        out.append(f"| {r} | " + " | ".join(_cell_text(v) for v in row) + " |")
    if ws.max_row > max_rows:
        out.append(f"\n_… còn {ws.max_row - max_rows} dòng nữa (tăng `max_rows` để xem)._")
    ranges = _merged_ranges(ws)
    if ranges:
        shown = ", ".join(f"`{r}`" for r in ranges[:_MERGED_PREVIEW])
        more = f" … và {len(ranges) - _MERGED_PREVIEW} vùng nữa" if len(ranges) > _MERGED_PREVIEW else ""
        out.append(
            f"\n_Ô gộp ({len(ranges)}): {shown}{more}. Giá trị nằm ở ô góc trên trái, "
            "các ô còn lại để trống._"
        )
    return "\n".join(out)

